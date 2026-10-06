"""Посещаемость группы: староста отмечает сам, видит отметки преподавателя, журнал и отметки в расписании."""

import io
from datetime import datetime

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select

from app import attendance, journal, users
from app.api.routes import create_app
from app.config import config
from app.importer import import_schedule, import_students
from app.models import ROLE_STAROSTA, AttendanceSession, Student, Teacher, User
from tests.conftest import PASSWORD

# Понедельник чётной недели, 10:50: физкультура (9:00) прошла, зоология (10:40) идёт
MONDAY = datetime(2026, 10, 5, 10, 50, tzinfo=config.tz)
DAY = "2026-10-05"


@pytest.fixture
async def c(database, schedule_xlsx, students_xlsx, monkeypatch):
    monkeypatch.setattr(attendance, "local_now", lambda: MONDAY)
    async with database.session() as s:
        await import_schedule(s, schedule_xlsx)
        await import_students(s, students_xlsx)
    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://t") as client:
        yield client


async def _account(database, login: str, fio: str, *, student: bool = True, role: str | None = None) -> dict:
    async with database.session() as s:
        user = await users.create_user(s, login, PASSWORD, fio)
        token = await users.issue_token(s, user)
        if student:
            user.student_id = await s.scalar(select(Student.id).where(Student.full_name == fio))
        if role:
            user.role = role
        await s.commit()
    return {"Authorization": f"Bearer {token}"}


async def _teacher(database) -> dict:
    h = await _account(database, "sidorov", "Сидоров Сергей Сергеевич", student=False)
    async with database.session() as s:
        uid = await s.scalar(select(User.id).where(User.login == "sidorov"))
        s.add(Teacher(user_id=uid, full_name="Сидоров Сергей Сергеевич"))
        await s.commit()
    return h


async def _monday(c, h) -> list[dict]:
    week = (await c.get("/api/schedule", params={"date": DAY}, headers=h)).json()
    return next(d for d in week["days"] if d["date"] == DAY)["lessons"]


async def _sid(database, fio: str) -> int:
    async with database.session() as s:
        return await s.scalar(select(Student.id).where(Student.full_name == fio))


async def test_starosta_marks_lessons_and_sees_them_in_schedule(c, database):
    star = await _account(database, "star", "Тестов Тест Тестович", role=ROLE_STAROSTA)
    anna = await _account(database, "anna", "Пробная Анна Сергеевна")
    anna_id = await _sid(database, "Пробная Анна Сергеевна")

    # Начавшиеся пары — отметить можно; у студента без отметок ничего не показываем
    lessons = await _monday(c, star)
    assert [(l["start"], l["attendance"]) for l in lessons] == [
        ("09:00", {"staff": True, "session": None}), ("10:40", {"staff": True, "session": None})]
    assert all("attendance" not in l for l in await _monday(c, anna))
    # В четверг пары ещё впереди
    week = (await c.get("/api/schedule", params={"date": DAY}, headers=star)).json()
    assert all("attendance" not in l for d in week["days"][1:] for l in d["lessons"])

    # Прошедшая физкультура — вручную, без кода
    r = await c.post("/api/admin/attendance", headers=star,
                     json={"group": "ГР 1", "date": DAY, "start": "09:00", "subject": "Физкультура"})
    assert r.status_code == 200, r.text
    x = r.json()
    assert x["editable"] is True and x["by_teacher"] is False and x["open"] is False and x["code"] is None
    assert x["teacher"] == "Петров П. П." and x["end"] == "10:30"
    assert [p["full_name"] for p in x["roster"]] == ["Пробная Анна Сергеевна", "Тестов Тест Тестович"]
    r = await c.put(f"/api/admin/attendance/{x['id']}/marks/{anna_id}", headers=star, json={"present": True})
    assert r.json()["present"] == 1 and r.json()["roster"][0]["method"] == "manual"
    # Повторно та же пара — та же отметка
    again = (await c.post("/api/admin/attendance", headers=star,
                          json={"group": "ГР 1", "date": DAY, "start": "09:00", "subject": "Физкультура"})).json()
    assert again["id"] == x["id"]

    # В расписании: старосте — сколько из группы, студенту — его отметка
    pe = (await _monday(c, star))[0]["attendance"]
    assert pe == {"staff": True, "session": x["id"], "present": 1, "total": 2, "open": False,
                  "by_teacher": False, "marked": False, "at": None}
    assert (await _monday(c, anna))[0]["attendance"]["marked"] is True

    # Идущая пара — с кодом: студенты отмечаются сами
    r = await c.post("/api/admin/attendance", headers=star,
                     json={"group": "ГР 1", "date": DAY, "start": "10:40", "subject": "Зоология", "code": True})
    z = r.json()
    assert len(z["code"]) == 4 and z["open"] is True
    r = await c.post("/api/attendance/checkin", headers=anna, json={"code": z["code"]})
    assert r.status_code == 200 and r.json()["subject"] == "Зоология" and r.json()["teacher"] == "Сидоров С. С."
    assert (await _monday(c, star))[1]["attendance"]["present"] == 1
    # Код у старосты один: новая отметка закрывает прошлую
    closed = (await c.post(f"/api/admin/attendance/{z['id']}/close", headers=star)).json()
    assert closed["open"] is False and closed["code"] is None
    reopened = (await c.post(f"/api/admin/attendance/{z['id']}/code", headers=star)).json()
    assert reopened["open"] is True and reopened["code"] and reopened["present"] == 1

    # Журнал группы: каждая пара и каждый студент
    j = (await c.get("/api/admin/attendance", params={"group": "ГР 1"}, headers=star)).json()
    assert [(b["subject"], b["present"], b["total"]) for b in j["sessions"]] == [("Зоология", 1, 2), ("Физкультура", 1, 2)]
    assert j["rate"] == 50
    assert {st["full_name"]: (st["attended"], st["total"], st["missed"]) for st in j["students"]} == {
        "Пробная Анна Сергеевна": (2, 2, []), "Тестов Тест Тестович": (0, 2, [z["id"], x["id"]])}

    # Excel: лист на предмет
    r = await c.get("/api/admin/attendance/export", params={"group": "ГР 1"}, headers=star)
    assert r.status_code == 200 and "filename*=UTF-8''" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert sorted(wb.sheetnames) == ["Зоология", "Физкультура"]
    rows = list(wb["Физкультура"].values)
    assert "05.10 09:00" in rows[0][2] and "Петров" in rows[0][2]
    assert rows[1:] == [("ГР 1", "Пробная Анна Сергеевна", "+", 1, 1, 100), ("ГР 1", "Тестов Тест Тестович", "—", 0, 1, 0)]

    # Удалить свою отметку можно
    assert (await c.delete(f"/api/admin/attendance/{x['id']}", headers=star)).json() == {"ok": True}
    assert (await _monday(c, star))[0]["attendance"] == {"staff": True, "session": None}


async def test_teacher_session_is_read_only_for_starosta(c, database):
    star = await _account(database, "star", "Тестов Тест Тестович", role=ROLE_STAROSTA)
    anna = await _account(database, "anna", "Пробная Анна Сергеевна")
    t = await _teacher(database)
    x = (await c.post("/api/teacher/sessions", headers=t, json={
        "groups": ["ГР 1"], "subject": "Зоология", "start": "10:40", "end": "12:10"})).json()
    await c.post("/api/attendance/checkin", headers=anna, json={"code": x["code"]})

    zoo = (await _monday(c, star))[1]["attendance"]
    assert zoo["session"] == x["id"] and zoo["by_teacher"] is True and zoo["present"] == 1 and zoo["open"] is True
    d = (await c.get(f"/api/admin/attendance/{x['id']}", headers=star)).json()
    # Код преподавателя старосте не показываем, отметки — только смотреть
    assert d["editable"] is False and d["code"] is None and d["present"] == 1
    anna_id = await _sid(database, "Пробная Анна Сергеевна")
    r = await c.put(f"/api/admin/attendance/{x['id']}/marks/{anna_id}", headers=star, json={"present": False})
    assert r.status_code == 403 and "преподаватель" in r.json()["detail"]
    assert (await c.delete(f"/api/admin/attendance/{x['id']}", headers=star)).status_code == 403
    assert (await c.post(f"/api/admin/attendance/{x['id']}/code", headers=star)).status_code == 403
    # Отметить ту же пару старосте — вернётся отметка преподавателя, без кода
    same = (await c.post("/api/admin/attendance", headers=star, json={
        "group": "ГР 1", "date": DAY, "start": "10:40", "subject": "Зоология", "code": True})).json()
    assert same["id"] == x["id"] and same["code"] is None

    # В расписании преподавателя — сколько отметились
    week = (await c.get("/api/schedule", params={"date": DAY}, headers=t)).json()
    assert week["days"][0]["lessons"][0]["attendance"] == {
        "teacher": True, "session": x["id"], "present": 1, "total": 2, "open": True}


async def test_journal_permissions_and_checks(c, database):
    star = await _account(database, "star", "Тестов Тест Тестович", role=ROLE_STAROSTA)
    anna = await _account(database, "anna", "Пробная Анна Сергеевна")
    other = await _account(database, "elkin", "Ёлкин Пётр Иванович", role=ROLE_STAROSTA)  # староста ГР 2
    body = {"group": "ГР 1", "date": DAY, "start": "09:00", "subject": "Физкультура"}

    assert (await c.get("/api/admin/attendance", params={"group": "ГР 1"}, headers=anna)).status_code == 403
    assert (await c.post("/api/admin/attendance", headers=anna, json=body)).status_code == 403
    assert (await c.post("/api/admin/attendance", headers=other, json=body)).status_code == 403
    x = (await c.post("/api/admin/attendance", headers=star, json=body)).json()
    assert (await c.get(f"/api/admin/attendance/{x['id']}", headers=other)).status_code == 404

    async def fails(change: dict, text: str, status: int = 400):
        r = await c.post("/api/admin/attendance", headers=star, json={**body, **change})
        assert r.status_code == status and text in r.json()["detail"], r.text

    await fails({"date": "2026-10-08", "start": "11:00", "subject": "Анатомия"}, "ещё не началась")
    await fails({"date": "2026-09-28", "code": True}, "только на сегодняшней")
    await fails({"start": "08:00"}, "Такой пары", 404)
    await fails({"date": "2026-09-29", "subject": "Физкультура"}, "Такой пары", 404)  # во вторник её нет
    # Прошлый понедельник — задним числом можно
    old = (await c.post("/api/admin/attendance", headers=star, json={**body, "date": "2026-09-28"})).json()
    assert old["date"] == "2026-09-28" and old["id"] != x["id"]


def _session(**kw) -> AttendanceSession:
    base = dict(subject="Пара", start_time=None, half=0, opened_by=None,
                created_at=datetime(2026, 10, 5, 7, 0), groups=[], marks=[])  # 10:00 по Москве
    return AttendanceSession(**{**base, **kw})


def test_match_sessions():
    items = [("09:00", "10:30", "Физкультура", None), ("10:40", "12:10", "Зоология", None),
             ("10:40", "12:10", "Химия", 2)]
    by_time = _session(subject="Что-то", start_time="10:40")
    by_name = _session(subject="  физкультура ")
    other_half = _session(subject="Химия", start_time="10:40", half=1)
    got = journal.match_sessions(items, [by_time, by_name, other_half])
    assert got == {0: by_name, 1: by_time}
    # Без времени и с другим названием — по тому, когда открыли: 10:35 — уже следующая пара
    by_clock = _session(subject="Своя пара", created_at=datetime(2026, 10, 5, 7, 35))
    assert journal.match_sessions(items[:2], [by_clock]) == {1: by_clock}
    assert journal.match_sessions(items, [_session(subject="Другое", created_at=datetime(2026, 10, 5, 12, 0))]) == {}
    # Пару отметили и преподаватель, и староста — показываем отметку преподавателя
    star = _session(start_time="09:00", opened_by=5, created_at=datetime(2026, 10, 5, 6, 0))
    teacher = _session(start_time="09:00", created_at=datetime(2026, 10, 5, 6, 10))
    assert journal.match_sessions(items, [star, teacher])[0] is teacher


async def test_student_sees_own_attendance(c, database):
    star = await _account(database, "star", "Тестов Тест Тестович", role=ROLE_STAROSTA)
    anna = await _account(database, "anna", "Пробная Анна Сергеевна")
    anna_id = await _sid(database, "Пробная Анна Сергеевна")
    empty = (await c.get("/api/attendance/stats", headers=anna)).json()
    assert empty == {"attended": 0, "total": 0, "rate": None, "subjects": [], "missed": []}

    for start, subject in (("09:00", "Физкультура"), ("10:40", "Зоология")):
        x = (await c.post("/api/admin/attendance", headers=star,
                          json={"group": "ГР 1", "date": DAY, "start": start, "subject": subject})).json()
        if subject == "Физкультура":
            await c.put(f"/api/admin/attendance/{x['id']}/marks/{anna_id}", headers=star, json={"present": True})
    old = (await c.post("/api/admin/attendance", headers=star,
                        json={"group": "ГР 1", "date": "2026-09-28", "start": "09:00", "subject": "Физкультура"})).json()
    assert old["present"] == 0

    st = (await c.get("/api/attendance/stats", headers=anna)).json()
    assert (st["attended"], st["total"], st["rate"]) == (1, 3, 33)
    assert st["subjects"] == [
        {"subject": "Зоология", "attended": 0, "total": 1, "rate": 0},
        {"subject": "Физкультура", "attended": 1, "total": 2, "rate": 50}]
    assert [(m["date"], m["subject"]) for m in st["missed"]] == [(DAY, "Зоология"), ("2026-09-28", "Физкультура")]
    # Без группы — нельзя
    nobody = await _account(database, "nobody", "Никто Никого Никакович", student=False)
    assert (await c.get("/api/attendance/stats", headers=nobody)).status_code == 403
