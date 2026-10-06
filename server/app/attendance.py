"""Отметка на паре по коду и преподаватели.

Преподаватель в начале пары создаёт код из 4 цифр — он работает минуту. Студент вводит код
в приложении и оказывается в списке присутствующих этой пары.
"""

import io
import logging
import re
import secrets
import time
from collections import Counter, defaultdict, deque
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app import db, notify, schedule
from app.security import INVITE_CODE_LEN, clean_code as clean_any_code, pretty_code, random_code
from app.config import config
from app.models import (
    CHANGE_ADD, CHANGE_EDIT, AttendanceGroup, AttendanceMark, AttendanceSession, Lesson, Student, Teacher,
    TeacherInvite, User,
)

log = logging.getLogger(__name__)

CODE_LEN = 4
CODE_TTL = 60                 # секунд работает код
SESSION_HOURS = 3             # незавершённая пара сама закрывается через столько часов
FAIL_LIMIT, FAIL_WINDOW = 5, 5 * 60   # неверных кодов подряд — и пауза
INVITE_DAYS = 7


class AttendanceError(Exception):
    """Отметиться или открыть пару нельзя; текст — для человека."""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def utcnow() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def local_now() -> datetime:
    return datetime.now(config.tz)


def local_hhmm(dt: datetime | None) -> str | None:
    if dt is None:
        return None
    return dt.replace(tzinfo=timezone.utc).astimezone(config.tz).strftime("%H:%M")


# --- приглашения преподавателей ------------------------------------------------

def clean_invite(text: str | None) -> str | None:
    """«abcd-efgh», «АВСD EFGH» → «ABCDEFGH»; None — не похоже на код приглашения."""
    return clean_any_code(text, INVITE_CODE_LEN)


def pretty_invite(code: str) -> str:
    return pretty_code(code)


async def create_invite(s: AsyncSession, owner: User) -> TeacherInvite:
    if not owner.is_owner:
        raise AttendanceError("Коды для преподавателей выдаёт только главный админ", 403)
    while True:
        code = random_code(INVITE_CODE_LEN)
        if not await s.get(TeacherInvite, code):
            break
    invite = TeacherInvite(code=code, created_by=owner.id, expires_at=utcnow() + timedelta(days=INVITE_DAYS))
    s.add(invite)
    await s.commit()
    log.info("Главный админ %s создал код для преподавателя", owner.id)
    return invite


async def valid_invite(s: AsyncSession, text: str | None) -> TeacherInvite | None:
    code = clean_invite(text)
    invite = await s.get(TeacherInvite, code) if code else None
    if invite is None or invite.expires_at <= utcnow():
        return None
    return invite


async def register_teacher(
    s: AsyncSession, user_id: int, code: str, full_name: str, schedule_name: str | None = None
) -> Teacher:
    """Тратит приглашение и создаёт преподавателя. Код проверяем ещё раз: его могли отменить."""
    invite = await valid_invite(s, code)
    if invite is None:
        raise AttendanceError("Код приглашения не подошёл или устарел. Попросите у главного админа новый")
    teacher = await s.get(Teacher, user_id)
    if teacher is None:
        teacher = Teacher(user_id=user_id, full_name=full_name)
        s.add(teacher)
    teacher.full_name = full_name
    teacher.schedule_name = schedule_name
    await s.delete(invite)
    await s.commit()
    log.info("Зарегистрирован преподаватель %s", user_id)
    return teacher


async def remove_teacher(s: AsyncSession, user_id: int) -> bool:
    teacher = await s.get(Teacher, user_id)
    if teacher is None:
        return False
    await s.delete(teacher)  # пары в журнале остаются, с именем преподавателя
    await s.commit()
    return True


# --- преподаватель в расписании ----------------------------------------------------
# В графе «Преподаватель» пишут по-разному: «Сидоров С. С.», «Сидоров С.С.», «доц. Сидоров С.С.»,
# полностью или несколько человек через запятую. Сравниваем фамилию и инициалы.

TITLES = {"доц", "доцент", "проф", "профессор", "ст", "старший", "преп", "преподаватель",
          "асс", "ассистент", "зав", "кафедрой", "декан"}


@dataclass(frozen=True)
class Person:
    surname: str
    initials: tuple[str, ...]


def person_key(text: str | None) -> Person | None:
    """«доц. Сидоров С.С.», «Сидоров Сергей Сергеевич», «С. С. Сидоров» → Person("сидоров", ("с", "с"))."""
    words = re.findall(r"[а-яa-z-]+", (text or "").casefold().replace("ё", "е"))
    while words and words[0] in TITLES:
        words.pop(0)
    i = next((n for n, w in enumerate(words) if len(w.strip("-")) > 1), None)
    if i is None:
        return None
    after = [w[0] for w in words[i + 1:i + 3]]
    before = [w for w in words[:i] if len(w) == 1][:2]
    return Person(words[i].strip("-"), tuple(after or before))


def same_person(a: Person, b: Person) -> bool:
    """Фамилия совпадает, а инициалы не противоречат (у кого-то их может не быть)."""
    return a.surname == b.surname and all(x == y for x, y in zip(a.initials, b.initials))


def people_in(field: str | None) -> list[tuple[str, Person]]:
    """Все преподаватели из графы расписания: (как записан, ключ)."""
    out = []
    for part in re.split(r"[,;/\n]+|\sи\s", field or ""):
        name = " ".join(part.split())
        key = person_key(name)
        if key:
            out.append((name, key))
    return out


def teacher_key(teacher: Teacher) -> Person | None:
    return person_key(teacher.schedule_name or teacher.full_name)


async def schedule_people(s: AsyncSession) -> list[dict]:
    """Все преподаватели из расписания: как записаны, сколько пар, какие предметы у каких подгрупп."""
    per: dict[Person, dict] = {}
    for field, subject, group in await s.execute(select(Lesson.teacher, Lesson.subject, Lesson.group_name)):
        for name, key in people_in(field):
            item = per.setdefault(key, {"names": Counter(), "lessons": 0, "subjects": defaultdict(set)})
            item["names"][name] += 1
            item["lessons"] += 1
            item["subjects"][subject].add(group)
    people = [
        {"name": item["names"].most_common(1)[0][0], "names": list(item["names"]), "key": key,
         "lessons": item["lessons"],
         "subjects": [{"subject": x, "groups": sorted(gs)} for x, gs in sorted(item["subjects"].items())]}
        for key, item in per.items()
    ]
    return sorted(people, key=lambda p: p["name"].casefold())


async def find_in_schedule(s: AsyncSession, full_name: str) -> list[dict]:
    """Однофамильцы из расписания; compatible — инициалы подходят к этому ФИО. Подходящие — первыми."""
    me = person_key(full_name)
    if me is None:
        return []
    found = [{**p, "compatible": same_person(me, p["key"])} for p in await schedule_people(s)
             if p["key"].surname == me.surname]
    return sorted(found, key=lambda p: not p["compatible"])


async def teacher_lessons(s: AsyncSession, teacher: Teacher) -> list[Lesson]:
    """Строки расписания, где преподаватель стоит в графе «Преподаватель»."""
    me = teacher_key(teacher)
    if me is None:
        return []
    lessons = (await s.scalars(select(Lesson))).all()
    return [l for l in lessons if any(same_person(me, key) for _, key in people_in(l.teacher))]


async def teacher_user_ids(s: AsyncSession, *teacher_fields: str | None) -> list[int]:
    """Аккаунты преподавателей, которые записаны в этих графах «Преподаватель»."""
    keys = [key for f in teacher_fields for _, key in people_in(f)]
    if not keys:
        return []
    out = []
    for t in (await s.scalars(select(Teacher))).all():
        me = teacher_key(t)
        if me is not None and any(same_person(me, k) for k in keys):
            out.append(t.user_id)
    return out


async def teacher_profile(s: AsyncSession, teacher: Teacher) -> dict:
    """Как он записан в расписании, его предметы и подгруппы — всё из расписания."""
    lessons = await teacher_lessons(s, teacher)
    me = teacher_key(teacher)
    names = Counter(name for l in lessons for name, key in people_in(l.teacher) if me and same_person(me, key))
    subjects: dict[str, set[str]] = defaultdict(set)
    for l in lessons:
        subjects[l.subject].add(l.group_name)
    return {
        "schedule_name": teacher.schedule_name or (names.most_common(1)[0][0] if names else None),
        "lessons": len(lessons),
        "subjects": [{"subject": x, "groups": sorted(gs)} for x, gs in sorted(subjects.items())],
        "groups": sorted({l.group_name for l in lessons}),
    }


async def teacher_groups(s: AsyncSession, teacher: Teacher) -> list[str]:
    """Подгруппы, для которых преподаватель может открыть отметку: у них есть его пары."""
    return sorted({l.group_name for l in await teacher_lessons(s, teacher)})


async def set_schedule_name(s: AsyncSession, teacher: Teacher, name: str | None) -> None:
    """Выбрать себя в расписании вручную (если там записан иначе); None — снова искать по ФИО."""
    name = " ".join((name or "").split()) or None
    if name is not None and not any(name in p["names"] for p in await schedule_people(s)):
        raise AttendanceError("Такого преподавателя нет в расписании")
    teacher.schedule_name = name
    await s.commit()


async def notify_owners(s: AsyncSession, title: str, body: str, skip: int | None = None) -> None:
    """Уведомление главным админам (без коммита)."""
    await notify.push(s, await notify.owner_ids(s), title, body, kind=notify.KIND_TEACHER,
                      skip=[skip] if skip else ())


# --- пара с отметкой -----------------------------------------------------------

def is_open(x: AttendanceSession, now: datetime | None = None) -> bool:
    now = now or utcnow()
    return x.closed_at is None and x.created_at > now - timedelta(hours=SESSION_HOURS)


def code_left(x: AttendanceSession, now: datetime | None = None) -> float:
    """Сколько секунд ещё работает код; 0 — не работает."""
    now = now or utcnow()
    if not is_open(x, now) or not x.code or not x.code_expires_at:
        return 0.0
    return max(0.0, (x.code_expires_at - now).total_seconds())


async def _new_code(s: AsyncSession) -> str:
    """Случайный код, которого нет среди действующих сейчас."""
    now = utcnow()
    busy = set((await s.scalars(
        select(AttendanceSession.code).where(AttendanceSession.code_expires_at > now)
    )).all())
    for _ in range(100):
        code = f"{secrets.randbelow(10 ** CODE_LEN):0{CODE_LEN}d}"
        if code not in busy:
            return code
    raise AttendanceError("Не получилось создать код — попробуйте ещё раз", 503)


async def open_session(
    s: AsyncSession, teacher: Teacher, groups: list[str], half: int, subject: str,
    start: str | None = None, end: str | None = None,
) -> AttendanceSession:
    allowed = set(await teacher_groups(s, teacher))
    if not allowed:
        raise AttendanceError("В расписании нет ваших пар. Проверьте на вкладке «Я учитель», как вы там записаны")
    groups = sorted(set(groups))
    if not groups:
        raise AttendanceError("Выберите хотя бы одну подгруппу")
    if any(g not in allowed for g in groups):
        raise AttendanceError("Код можно создать только для подгрупп, у которых вы ведёте пары", 403)
    if half not in (0, 1, 2):
        raise AttendanceError("Половина группы — 1 или 2")
    now = utcnow()
    # Прошлая пара преподавателя закрывается: одновременно открыта только одна
    for old in (await s.scalars(select(AttendanceSession).where(
        AttendanceSession.teacher_id == teacher.user_id, AttendanceSession.closed_at.is_(None)
    ))).all():
        old.closed_at = now
    x = AttendanceSession(
        teacher_id=teacher.user_id, teacher_name=teacher.full_name,
        subject=" ".join((subject or "").split())[:300] or "Пара",
        lesson_date=local_now().date(), start_time=start, end_time=end, half=half,
        code=await _new_code(s), code_expires_at=now + timedelta(seconds=CODE_TTL), created_at=now,
        groups=[AttendanceGroup(group_name=g) for g in groups], marks=[],
    )
    s.add(x)
    await s.commit()
    log.info("Преподаватель %s открыл отметку для %s", teacher.user_id, ", ".join(groups))
    return x


async def new_code(s: AsyncSession, x: AttendanceSession) -> AttendanceSession:
    if not is_open(x):
        raise AttendanceError("Отметка на этой паре уже завершена")
    x.code = await _new_code(s)
    x.code_expires_at = utcnow() + timedelta(seconds=CODE_TTL)
    await s.commit()
    return x


async def close_session(s: AsyncSession, x: AttendanceSession) -> None:
    now = utcnow()
    if x.closed_at is None:
        x.closed_at = now
    if x.code_expires_at and x.code_expires_at > now:
        x.code_expires_at = now
    await s.commit()


async def teacher_session(s: AsyncSession, teacher: Teacher, session_id: int) -> AttendanceSession:
    x = await s.get(AttendanceSession, session_id)
    if x is None or x.teacher_id != teacher.user_id:
        raise AttendanceError("Пара не найдена", 404)
    return x


async def active_session(s: AsyncSession, teacher: Teacher) -> AttendanceSession | None:
    x = await s.scalar(
        select(AttendanceSession)
        .where(AttendanceSession.teacher_id == teacher.user_id, AttendanceSession.closed_at.is_(None))
        .order_by(AttendanceSession.created_at.desc())
    )
    return x if x is not None and is_open(x) else None


# --- кто должен быть на паре ---------------------------------------------------

class Rosters:
    """Студенты подгрупп и их половины — один раз на много пар (для журнала и статистики)."""

    def __init__(self, students: list[Student], halves: dict[int, int | None], linked: set[int]):
        self.by_group: dict[str, list[Student]] = defaultdict(list)
        for st in students:
            self.by_group[st.group_name].append(st)
        self.by_id = {st.id: st for st in students}
        self.halves = halves
        self.linked = linked

    @classmethod
    async def load(cls, s: AsyncSession, groups: set[str], extra_ids: set[int] = frozenset()) -> "Rosters":
        stmt = select(Student).where(Student.group_name.in_(groups))
        students = list((await s.scalars(stmt)).all()) if groups else []
        have = {st.id for st in students}
        missing = set(extra_ids) - have
        if missing:  # отмеченные, кого с тех пор перевели в другую группу
            students += list((await s.scalars(select(Student).where(Student.id.in_(missing)))).all())
        ids = [st.id for st in students]
        users = (await s.execute(
            select(User.student_id, User.half).where(User.student_id.in_(ids))
        )).all() if ids else []
        students.sort(key=lambda st: (st.group_name, st.full_name))
        return cls(students, {sid: half for sid, half in users}, {sid for sid, _ in users})

    def expected(self, x: AttendanceSession) -> list[Student]:
        """Кто должен быть на паре: студенты её подгрупп (на паре половины — без другой половины)
        и все, кто отметился."""
        marked = {m.student_id for m in x.marks}
        out = []
        for g in x.group_names:
            for st in self.by_group.get(g, []):
                half = self.halves.get(st.id)
                if x.half and half and half != x.half and st.id not in marked:
                    continue
                out.append(st)
        seen = {st.id for st in out}
        out += [self.by_id[sid] for sid in sorted(marked - seen) if sid in self.by_id]
        return out


async def rosters_for(s: AsyncSession, sessions: list[AttendanceSession]) -> Rosters:
    groups = {g for x in sessions for g in x.group_names}
    marked = {m.student_id for x in sessions for m in x.marks}
    return await Rosters.load(s, groups, marked)


def session_brief(x: AttendanceSession, r: Rosters) -> dict:
    now = utcnow()
    return {
        "id": x.id,
        "subject": x.subject,
        "groups": x.group_names,
        "half": x.half,
        "date": x.lesson_date.isoformat(),
        "start": x.start_time,
        "end": x.end_time,
        "started_at": local_hhmm(x.created_at),
        "teacher": x.teacher_name,
        "open": is_open(x, now),
        "present": len(x.marks),
        "total": len(r.expected(x)),
    }


async def session_detail(s: AsyncSession, x: AttendanceSession) -> dict:
    r = await rosters_for(s, [x])
    now = utcnow()
    left = code_left(x, now)
    marks = {m.student_id: m for m in x.marks}
    roster = []
    for st in r.expected(x):
        m = marks.get(st.id)
        roster.append({
            "id": st.id, "full_name": st.full_name, "group": st.group_name,
            "present": m is not None, "at": local_hhmm(m.marked_at) if m else None,
            "method": m.method if m else None, "in_app": st.id in r.linked,
        })
    return {
        **session_brief(x, r),
        "code": x.code if left > 0 else None,
        "expires_in": round(left, 1),
        "ttl": CODE_TTL,
        "roster": roster,
    }


async def set_mark(s: AsyncSession, x: AttendanceSession, student_id: int, present: bool) -> None:
    """Отметка вручную: студент без телефона или ошибся с кодом."""
    r = await rosters_for(s, [x])
    if student_id not in {st.id for st in r.expected(x)}:
        raise AttendanceError("Этого студента нет в подгруппах пары", 404)
    mark = await s.get(AttendanceMark, (x.id, student_id))
    if present and mark is None:
        s.add(AttendanceMark(session_id=x.id, student_id=student_id, marked_at=utcnow(), method="manual"))
    elif not present and mark is not None:
        await s.delete(mark)
    await s.commit()
    await s.refresh(x, attribute_names=["marks"])


# --- студент отмечается --------------------------------------------------------

_fails: dict[int, deque] = defaultdict(deque)


def _recent_fails(user_id: int) -> deque:
    q = _fails[user_id]
    now = time.monotonic()
    while q and now - q[0] > FAIL_WINDOW:
        q.popleft()
    return q


def clean_code(text: str | None) -> str | None:
    code = "".join((text or "").split())
    return code if len(code) == CODE_LEN and code.isdigit() else None


async def checkin(s: AsyncSession, user: User, text: str | None) -> tuple[AttendanceSession, AttendanceMark, bool]:
    """Отмечает студента по коду. Возвращает (пара, отметка, отмечался ли уже раньше)."""
    if user.student is None:
        raise AttendanceError("Отмечаться могут студенты, которые уже в группе", 403)
    fails = _recent_fails(user.id)
    if len(fails) >= FAIL_LIMIT:
        wait = max(1, round((FAIL_WINDOW - (time.monotonic() - fails[0])) / 60))
        raise AttendanceError(f"Слишком много неверных кодов. Попробуй через {wait} мин", 429)
    code = clean_code(text)
    if code is None:
        raise AttendanceError("Код — это 4 цифры", 422)
    now = utcnow()
    found = (await s.scalars(select(AttendanceSession).where(
        AttendanceSession.code == code,
        AttendanceSession.code_expires_at > now,
        AttendanceSession.closed_at.is_(None),
    ))).all()
    group = user.group_name
    x = next((f for f in found if group in f.group_names), None)
    if x is None:
        fails.append(time.monotonic())
        if found:
            raise AttendanceError("Этот код для другой группы")
        raise AttendanceError("Код неверный или уже не действует. Попроси преподавателя показать новый")
    if x.half and user.half and user.half != x.half:
        raise AttendanceError(f"Этот код для {x.half}-й половины группы, а у тебя в профиле — {user.half}-я")
    mark = await s.get(AttendanceMark, (x.id, user.student_id))
    if mark is not None:
        return x, mark, True
    mark = AttendanceMark(session_id=x.id, student_id=user.student_id, marked_at=now, method="code")
    s.add(mark)
    try:
        await s.commit()
    except IntegrityError:  # нажал дважды — вторая отметка не нужна
        await s.rollback()
        return x, await s.get(AttendanceMark, (x.id, user.student_id)), True
    fails.clear()
    log.info("Студент %s отметился на паре %s", user.student_id, x.id)
    return x, mark, False


async def student_today(s: AsyncSession, user: User) -> list[dict]:
    """Пары с отметкой сегодня у группы студента и отметился ли он."""
    if user.student is None:
        return []
    sessions = (await s.scalars(
        select(AttendanceSession).join(AttendanceGroup)
        .where(AttendanceGroup.group_name == user.group_name, AttendanceSession.lesson_date == local_now().date())
        .order_by(AttendanceSession.created_at)
    )).all()
    now = utcnow()
    out = []
    for x in sessions:
        mark = next((m for m in x.marks if m.student_id == user.student_id), None)
        if x.half and user.half and x.half != user.half and mark is None:
            continue
        out.append({
            "id": x.id, "subject": x.subject, "start": x.start_time, "end": x.end_time,
            "teacher": x.teacher_name, "started_at": local_hhmm(x.created_at),
            "code_active": code_left(x, now) > 0, "marked": mark is not None,
            "marked_at": local_hhmm(mark.marked_at) if mark else None,
        })
    return out


# --- журнал преподавателя --------------------------------------------------------

async def teacher_sessions(s: AsyncSession, teacher: Teacher, subject: str | None = None) -> list[AttendanceSession]:
    stmt = (select(AttendanceSession).where(AttendanceSession.teacher_id == teacher.user_id)
            .order_by(AttendanceSession.lesson_date.desc(), AttendanceSession.created_at.desc()))
    if subject is not None:
        stmt = stmt.where(AttendanceSession.subject == subject)
    return list((await s.scalars(stmt)).all())


def _rate(part: int, total: int) -> int | None:
    return round(part * 100 / total) if total else None


async def teacher_summary(s: AsyncSession, teacher: Teacher, recent: int = 5) -> dict:
    sessions = await teacher_sessions(s, teacher)
    r = await rosters_for(s, sessions)
    profile = await teacher_profile(s, teacher)
    by_subject = {x["subject"]: x["groups"] for x in profile["subjects"]}
    group_r = await Rosters.load(s, set(profile["groups"]))
    per_subject = {x: [0, 0, 0] for x in by_subject}  # пар, было, должно было быть
    present = total = 0
    for x in sessions:
        marked, expected = len(x.marks), len(r.expected(x))
        present += marked
        total += expected
        if x.subject in per_subject:
            per_subject[x.subject][0] += 1
            per_subject[x.subject][1] += marked
            per_subject[x.subject][2] += expected
    return {
        "full_name": teacher.full_name,
        "schedule_name": profile["schedule_name"],
        "lessons": profile["lessons"],
        "subjects": [
            {"subject": x, "groups": by_subject[x], "sessions": n, "rate": _rate(was, should),
             "students": sum(len(group_r.by_group.get(g, [])) for g in by_subject[x])}
            for x, (n, was, should) in per_subject.items()
        ],
        "groups": profile["groups"],
        "sessions_total": len(sessions),
        "rate": _rate(present, total),
        "recent": [session_brief(x, r) for x in sessions[:recent]],
    }


async def subject_attendance(s: AsyncSession, teacher: Teacher, subject: str) -> dict:
    """Посещаемость по предмету на парах этого преподавателя — по каждому студенту.

    Подгруппы — те, с кем уже были пары; пока пар не было — по расписанию.
    """
    sessions = await teacher_sessions(s, teacher, subject)
    scheduled = {l.group_name for l in await teacher_lessons(s, teacher) if l.subject == subject}
    if not sessions and not scheduled:
        raise AttendanceError("Это не ваш предмет", 403)
    groups = {g for x in sessions for g in x.group_names} or scheduled
    r = await rosters_for(s, sessions)
    group_r = await Rosters.load(s, groups)
    counts: dict[int, list[int]] = {st.id: [0, 0] for st in group_r.by_id.values()}
    for x in sessions:
        marked = {m.student_id for m in x.marks}
        for st in r.expected(x):
            if st.id in counts:
                counts[st.id][1] += 1
                counts[st.id][0] += st.id in marked
    students = [
        {"id": st.id, "full_name": st.full_name, "group": st.group_name, "attended": counts[st.id][0],
         "total": counts[st.id][1], "rate": _rate(*counts[st.id]), "in_app": st.id in group_r.linked}
        for st in sorted(group_r.by_id.values(), key=lambda st: (st.group_name, st.full_name))
    ]
    was = sum(c[0] for c in counts.values())
    should = sum(c[1] for c in counts.values())
    return {"subject": subject, "sessions": len(sessions), "rate": _rate(was, should), "students": students}


def _sheet_title(name: str, used: set[str]) -> str:
    title = "".join(ch for ch in name if ch not in "[]:*?/\\")[:28].strip() or "Пары"
    base, n = title, 2
    while title in used:
        title, n = f"{base[:26]} {n}", n + 1
    used.add(title)
    return title


async def export_xlsx(s: AsyncSession, teacher: Teacher) -> bytes | None:
    """Журнал в Excel: лист на предмет, строки — студенты по подгруппам, столбцы — пары.
    None — пар ещё нет."""
    sessions = list(reversed(await teacher_sessions(s, teacher)))
    if not sessions:
        return None
    r = await rosters_for(s, sessions)
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for subject in sorted({x.subject for x in sessions}):
        ws = wb.create_sheet(_sheet_title(subject, used))
        mine = [x for x in sessions if x.subject == subject]
        ws.append(["Группа", "ФИО"] + [
            f"{x.lesson_date:%d.%m} {x.start_time or local_hhmm(x.created_at)}\n{', '.join(x.group_names)}"
            for x in mine
        ] + ["Был", "Из", "%"])
        expected = {x.id: {st.id for st in r.expected(x)} for x in mine}
        marked = {x.id: {m.student_id for m in x.marks} for x in mine}
        students = {sid for ids in expected.values() for sid in ids}
        for st in sorted((r.by_id[sid] for sid in students), key=lambda st: (st.group_name, st.full_name)):
            cells, was, should = [], 0, 0
            for x in mine:
                if st.id in marked[x.id]:
                    cells.append("+")
                    was += 1
                    should += 1
                elif st.id in expected[x.id]:
                    cells.append("—")
                    should += 1
                else:
                    cells.append("")  # на этой паре была другая подгруппа или половина
            ws.append([st.group_name, st.full_name, *cells, was, should, _rate(was, should)])
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 36
        ws.freeze_panes = "C2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2, min_col=3):
            for cell in row:
                cell.alignment = Alignment(horizontal="center")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def lesson_key(lesson: Lesson) -> tuple:
    return (lesson.subject, lesson.start_time, lesson.end_time, lesson.half or 0, lesson.kind)


def teacher_items(slots: list[schedule.Slot]) -> list[dict]:
    """Пары преподавателя за день. Одинаковые пары разных подгрупп (общая лекция) — одной строкой."""
    merged: dict[tuple, dict] = {}
    for l in slots:
        item = merged.setdefault(lesson_key(l), {
            "subject": l.subject, "kind": l.kind, "room": l.room, "start": l.start_time,
            "end": l.end_time, "half": l.half or 0, "pair_num": l.pair_num, "groups": [],
            "teacher": l.teacher, "status": l.status, "note": l.note, "moved": l.moved,
            "was": schedule.LessonDTO.of(l).was,
        })
        item["groups"].append(l.group_name)
    for item in merged.values():
        item["groups"].sort()
    return sorted(merged.values(), key=lambda l: (l["start"], l["subject"]))


def _teaches(me: Person, teacher_field: str | None) -> bool:
    return any(same_person(me, key) for _, key in people_in(teacher_field))


async def teacher_days(s: AsyncSession, teacher: Teacher, dates: list[date]) -> dict[date, schedule.Day]:
    """Пары преподавателя по дням — с отменами, заменами аудиторий и разовыми парами.
    Замена преподавателя на один день: пара уходит к тому, кто её ведёт в этот день."""
    me = teacher_key(teacher)
    if me is None:
        return {d: schedule.Day([], []) for d in dates}
    mine = await teacher_lessons(s, teacher)
    ids = {l.id for l in mine}
    changes = await schedule.changes_between(s, min(dates), max(dates))
    # Чужие пары, которые ему отдали на один день
    sub_ids = {c.lesson_id for c in changes
               if c.action == CHANGE_EDIT and c.lesson_id not in ids and c.teacher and _teaches(me, c.teacher)}
    lessons = mine + (list((await s.scalars(select(Lesson).where(Lesson.id.in_(sub_ids)))).all()) if sub_ids else [])
    relevant = [c for c in changes if c.lesson_id in ids or c.lesson_id in sub_ids
                or (c.action == CHANGE_ADD and _teaches(me, c.teacher))]
    start = await db.get_semester_start(s)
    out = {}
    for d in dates:
        day = schedule.build_day(lessons, relevant, d, start, None)
        out[d] = schedule.Day([l for l in day.lessons if _teaches(me, l.teacher)],
                              [l for l in day.cancelled if _teaches(me, l.teacher)])
    return out


async def teacher_today(s: AsyncSession, teacher: Teacher) -> dict:
    """Пары преподавателя сегодня (где он стоит в расписании) — чтобы открыть отметку одним нажатием."""
    mine = await teacher_lessons(s, teacher)
    d = local_now().date()
    lessons = teacher_items((await teacher_days(s, teacher, [d]))[d].lessons)
    x = await active_session(s, teacher)
    return {
        "groups": sorted({l.group_name for l in mine}),
        "lessons": lessons,
        "active": await session_detail(s, x) if x else None,
    }


# --- расписание преподавателя в чате ---------------------------------------------

SHORT_GROUP_RE = re.compile(r"^(.*?)\s*(\d+)\s*подгруппа$", re.IGNORECASE)


def short_group(group: str) -> str:
    """«БИА 2 ПОДГРУППА» → «БИА 2»."""
    m = SHORT_GROUP_RE.match(group or "")
    return f"{m.group(1)} {m.group(2)}" if m else group


def groups_line(item: dict) -> str:
    text = ", ".join(short_group(g) for g in item["groups"])
    return f"{text} · {item['half']}-я половина" if item["half"] else text
