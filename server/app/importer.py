"""Импорт расписания и списка студентов из xlsx."""

import io
import re
from dataclasses import dataclass, field
from datetime import datetime, time

from openpyxl import load_workbook
from sqlalchemy import delete, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app import db, notify
from app.models import WEEK_EVEN, WEEK_EVERY, WEEK_ODD, Lesson, Student, User

KIND_SCHEDULE = "schedule"
KIND_STUDENTS = "students"


class ImportError_(ValueError):
    pass


def name_key(full_name: str) -> str:
    """Ключ для сравнения ФИО: регистр, «ё» и лишние пробелы не важны."""
    return " ".join(full_name.replace("ё", "е").replace("Ё", "Е").split()).casefold()


def clean_group(value) -> str:
    return " ".join(str(value or "").split())


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def clean_email(text: str | None) -> str | None:
    """Почта без пробелов; None — пусто. ValueError — не похоже на почту."""
    email = (text or "").strip()
    if not email:
        return None
    if len(email) > 200 or not EMAIL_RE.match(email):
        raise ValueError(f"не похоже на почту: «{email}»")
    return email


HEADER_ALIASES = {
    "group": ["группа"],
    "day": ["день", "день недели"],
    "week": ["неделя", "чётность", "четность"],
    "pair": ["№ пары", "пара", "номер пары", "№"],
    "start": ["начало"],
    "end": ["конец", "окончание"],
    "subject": ["предмет", "дисциплина"],
    "kind": ["тип", "вид", "тип занятия"],
    "room": ["аудитория", "ауд", "ауд.", "кабинет"],
    "teacher": ["преподаватель", "препод", "преподаватель(и)"],
    "half": ["половина", "половина группы"],
    "fio": ["фио", "ф.и.о.", "студент"],
    "email": ["почта", "email", "e-mail", "эл. почта", "электронная почта", "корпоративная почта"],
}

DAYS = {
    "понедельник": 0, "пн": 0, "вторник": 1, "вт": 1, "среда": 2, "ср": 2,
    "четверг": 3, "чт": 3, "пятница": 4, "пт": 4, "суббота": 5, "сб": 5,
    "воскресенье": 6, "вс": 6,
}

HALF_IN_SUBJECT = re.compile(r"\s*\(\s*(\d)\s*-?\s*я\s+половина[^)]*\)", re.IGNORECASE)


def parse_day(value) -> int:
    key = _cell_str(value).lower().rstrip(".")
    if key not in DAYS:
        raise ImportError_(f"не понял день недели «{value}»")
    return DAYS[key]


def parse_week(value) -> str:
    v = _cell_str(value).lower().replace("ё", "е")
    if not v or v.startswith("кажд") or v in ("все", "обе", "-"):
        return WEEK_EVERY
    if v.startswith("нечет") or v.startswith("числ"):
        return WEEK_ODD
    if v.startswith("чет") or v.startswith("знам"):
        return WEEK_EVEN
    raise ImportError_(f"не понял неделю «{value}» (нужно: каждая / нечётная / чётная)")


def parse_time(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}"
    if isinstance(value, (int, float)) and 0 <= value < 1:  # доля суток в Excel
        minutes = round(value * 24 * 60)
        return f"{minutes // 60:02d}:{minutes % 60:02d}"
    m = re.fullmatch(r"(\d{1,2})[:.](\d{2})(?::\d{2})?", _cell_str(value))
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        raise ImportError_(f"не понял время «{value}» (нужно ЧЧ:ММ)")
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def parse_half(value) -> int | None:
    v = _cell_str(value)
    if not v:
        return None
    m = re.match(r"([12])", v)
    if not m:
        raise ImportError_(f"не понял половину группы «{value}» (нужно 1 или 2)")
    return int(m.group(1))


def split_half_from_subject(subject: str) -> tuple[str, int | None]:
    """«Предмет (1-я половина группы — уточни…)» → («Предмет», 1)."""
    m = HALF_IN_SUBJECT.search(subject)
    if not m:
        return subject, None
    return HALF_IN_SUBJECT.sub("", subject).strip(), int(m.group(1))


def _map_headers(row) -> dict[str, int]:
    lookup = {alias: key for key, aliases in HEADER_ALIASES.items() for alias in aliases}
    found: dict[str, int] = {}
    for idx, cell in enumerate(row):
        key = lookup.get(_cell_str(cell).lower())
        if key and key not in found:
            found[key] = idx
    return found


def _find_table(data: bytes) -> tuple[str, dict[str, int], list[tuple[int, tuple]]]:
    """Ищет лист с заголовками и определяет, что это: расписание или список студентов."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — openpyxl кидает разные ошибки на битые файлы
        raise ImportError_("не удалось открыть файл — нужен .xlsx") from e
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            for header_idx, row in enumerate(rows[:10]):
                headers = _map_headers(row)
                if {"group", "subject", "day"} <= headers.keys():
                    kind = KIND_SCHEDULE
                elif {"fio", "group"} <= headers.keys():
                    kind = KIND_STUDENTS
                else:
                    continue
                body = [
                    (header_idx + 2 + i, r)
                    for i, r in enumerate(rows[header_idx + 1:])
                    if any(c not in (None, "") for c in r)
                ]
                return kind, headers, body
    finally:
        wb.close()
    raise ImportError_(
        "не нашёл заголовки. Для расписания нужны колонки «Группа», «День», «Предмет», "
        "для списка студентов — «ФИО» и «Группа»"
    )


def detect_kind(data: bytes) -> str:
    return _find_table(data)[0]


@dataclass
class ParsedSchedule:
    lessons: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def parse_schedule(data: bytes, bells: list[tuple[str, str]]) -> ParsedSchedule:
    kind, h, body = _find_table(data)
    if kind != KIND_SCHEDULE:
        raise ImportError_("это список студентов, а не расписание")
    result = ParsedSchedule()

    def get(row, key):
        idx = h.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    for line, row in body:
        try:
            group = clean_group(get(row, "group"))
            subject = _cell_str(get(row, "subject"))
            if not group or not subject:
                raise ImportError_("пустая группа или предмет")
            subject, half = split_half_from_subject(subject)
            if h.get("half") is not None and _cell_str(get(row, "half")):
                half = parse_half(get(row, "half"))
            pair_raw = _cell_str(get(row, "pair"))
            pair_num = int(pair_raw) if pair_raw.isdigit() else None
            start = parse_time(get(row, "start"))
            end = parse_time(get(row, "end"))
            if (not start or not end) and pair_num and 1 <= pair_num <= len(bells):
                start = start or bells[pair_num - 1][0]
                end = end or bells[pair_num - 1][1]
            if not start or not end:
                raise ImportError_("нет времени и номера пары из сетки звонков")
            result.lessons.append(
                dict(
                    group_name=group,
                    weekday=parse_day(get(row, "day")),
                    week=parse_week(get(row, "week")),
                    pair_num=pair_num,
                    start_time=start,
                    end_time=end,
                    subject=subject,
                    kind=_cell_str(get(row, "kind")).lower(),
                    room=_cell_str(get(row, "room")),
                    teacher=_cell_str(get(row, "teacher")),
                    half=half,
                )
            )
        except ImportError_ as e:
            result.errors.append(f"строка {line}: {e}")
    return result


def parse_students(data: bytes) -> tuple[list[tuple[str, str, str | None]], list[str]]:
    """Строки (ФИО, группа, почта). Почта None — столбца нет или ячейка пустая."""
    kind, h, body = _find_table(data)
    if kind != KIND_STUDENTS:
        raise ImportError_("это расписание, а не список студентов")
    students: list[tuple[str, str, str | None]] = []
    errors: list[str] = []
    seen: dict[str, int] = {}
    for line, row in body:
        fio = _cell_str(row[h["fio"]] if h["fio"] < len(row) else None)
        group = clean_group(row[h["group"]] if h["group"] < len(row) else None)
        if not fio or not group:
            errors.append(f"строка {line}: пустое ФИО или группа")
            continue
        key = name_key(fio)
        if key in seen:
            errors.append(f"строка {line}: «{fio}» уже есть в строке {seen[key]}, пропускаю")
            continue
        seen[key] = line
        email = None
        if "email" in h and h["email"] < len(row):
            try:
                email = clean_email(_cell_str(row[h["email"]]))
            except ValueError as e:
                errors.append(f"строка {line}: {e}, почту пропускаю")
        students.append((fio, group, email))
    return students, errors


# --- запись в БД -----------------------------------------------------------

async def import_schedule(s: AsyncSession, data: bytes) -> str:
    parsed = parse_schedule(data, await db.get_bells(s))
    if not parsed.lessons:
        raise ImportError_("в файле нет ни одной пары\n" + "\n".join(parsed.errors[:10]))
    groups: dict[str, int] = {}
    for row in parsed.lessons:
        groups[row["group_name"]] = groups.get(row["group_name"], 0) + 1
    await s.execute(delete(Lesson).where(Lesson.group_name.in_(groups)))
    s.add_all(Lesson(**row) for row in parsed.lessons)
    for group in groups:
        await notify.push(s, await notify.group_member_ids(s, [group]), "🗓 Расписание обновлено",
                          f"Админ загрузил новое расписание группы {group}. Загляни во вкладку «Расписание».",
                          kind=notify.KIND_SCHEDULE)
    await s.commit()
    lines = [f"Расписание загружено: {len(parsed.lessons)} пар, групп — {len(groups)}"]
    lines += [f"• {g}: {n}" for g, n in sorted(groups.items())]
    return _with_errors(lines, parsed.errors)


async def import_students(s: AsyncSession, data: bytes) -> str:
    """Список студентов из xlsx. Записи, которых нет в файле, удаляются — кроме тех, кто уже
    вошёл в приложение и привязан к записи: их староста или админ удаляет вручную."""
    rows, errors = parse_students(data)
    if not rows:
        raise ImportError_("в файле нет ни одного студента\n" + "\n".join(errors[:10]))
    before = list((await s.scalars(select(Student).order_by(Student.id))).all())
    existing: dict[str, Student] = {}
    for st in before:
        existing.setdefault(st.name_key, st)
    linked = set((await s.scalars(select(User.student_id).where(User.student_id.is_not(None)))).all())
    added = updated = 0
    keep: set[int] = set()
    for fio, group, email in rows:
        key = name_key(fio)
        st = existing.get(key)
        if st is None:
            st = Student(full_name=fio, name_key=key, group_name=group, email=email)
            s.add(st)
            added += 1
        else:
            # Пустая ячейка почты не стирает ту, что уже есть
            new_email = email or st.email
            if st.group_name != group or st.full_name != fio or st.email != new_email:
                if st.group_name != group and st.id in linked:
                    # У другой группы своё деление на половины — выбор сбрасываем
                    await s.execute(update(User).where(User.student_id == st.id).values(half=None))
                st.group_name, st.full_name, st.email = group, fio, new_email
                updated += 1
            keep.add(st.id)
    stale = {st.id for st in before} - keep
    removed_ids = [sid for sid in stale if sid not in linked]
    kept_linked = len(stale) - len(removed_ids)
    if removed_ids:
        await s.execute(delete(Student).where(Student.id.in_(removed_ids)))
    await s.commit()
    lines = [
        f"Список студентов загружен: всего {len(rows)}",
        f"• новых: {added}, изменено: {updated}, удалено: {len(removed_ids)}",
    ]
    if kept_linked:
        lines.append(f"• нет в файле, но уже в приложении — оставлены: {kept_linked}")
    return _with_errors(lines, errors)


async def import_any(s: AsyncSession, data: bytes) -> str:
    if detect_kind(data) == KIND_SCHEDULE:
        return await import_schedule(s, data)
    return await import_students(s, data)


def _with_errors(lines: list[str], errors: list[str]) -> str:
    if errors:
        lines.append(f"\n⚠️ Пропущено строк: {len(errors)}")
        lines += errors[:15]
        if len(errors) > 15:
            lines.append(f"…и ещё {len(errors) - 15}")
    return "\n".join(lines)
